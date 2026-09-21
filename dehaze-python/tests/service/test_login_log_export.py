"""登录日志导出（F-AM-010 导出能力）单元测试。"""

import io
from datetime import datetime

from openpyxl import load_workbook

from app.service.auth_service import auth_service
from tests.stubs.factories import make_user_context


def _admin_user():
    return make_user_context(id=1, username="admin", roles=["ADMIN"])


def _normal_user():
    return make_user_context(id=2, username="normal")


def _make_doc(oid="id-1", user_id=1, username="admin", status=1, device_type="web"):
    return {
        "_id": oid,
        "user_id": user_id,
        "username": username,
        "ip": "1.2.3.4",
        "location": "",
        "browser": "Chrome 120",
        "os": "Windows",
        "device_type": device_type,
        "status": status,
        "message": "登录成功" if status == 1 else "密码错误",
        "create_time": datetime(2025, 1, 1, 12, 0, 0),
    }


def _load_workbook(content: bytes):
    return load_workbook(io.BytesIO(content))


async def test_export_contains_header_and_rows(mongo_db):
    await mongo_db["login_log"].insert_many(
        [
            _make_doc(oid="a1", status=1),
            _make_doc(oid="a2", username="normal", user_id=2, status=0, device_type="android"),
        ]
    )
    content = await auth_service.export_login_logs(user=_admin_user())
    ws = _load_workbook(content).active
    assert ws is not None

    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == (
        "用户名",
        "IP",
        "登录时间",
        "状态",
        "提示信息",
        "设备类型",
        "浏览器",
        "操作系统",
    )
    assert len(rows) == 3
    success_row = next(r for r in rows[1:] if r[0] == "admin")
    assert success_row[3] == "成功"
    assert success_row[5] == "web"
    failed_row = next(r for r in rows[1:] if r[0] == "normal")
    assert failed_row[3] == "失败"
    assert failed_row[5] == "android"


async def test_export_respects_filter_and_device_type(mongo_db):
    await mongo_db["login_log"].insert_many(
        [
            _make_doc(oid="a1", status=1),
            _make_doc(oid="a2", username="normal", user_id=2, status=0, device_type="android"),
        ]
    )
    content = await auth_service.export_login_logs(device_type="android", user=_admin_user())
    ws = _load_workbook(content).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][0] == "normal"


async def test_export_scoped_to_own_logs_for_normal_user(mongo_db):
    await mongo_db["login_log"].insert_many(
        [
            _make_doc(oid="a1", username="admin", user_id=1),
            _make_doc(oid="a2", username="normal", user_id=2),
        ]
    )
    content = await auth_service.export_login_logs(user=_normal_user())
    ws = _load_workbook(content).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][0] == "normal"
